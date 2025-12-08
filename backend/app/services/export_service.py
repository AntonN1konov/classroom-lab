import os
from typing import List, Dict
from datetime import datetime
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY

class ExportService:
    def __init__(self):
        self.exports_dir = "exports"
        os.makedirs(self.exports_dir, exist_ok=True)
    
    def export_session(
        self,
        session_name: str,
        messages: List[Dict],
        format: str = "pdf",
        include_context: bool = True
    ) -> str:
        """
        Экспорт сессии в документ
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = "".join(c for c in session_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
        filename = f"{safe_name}_{timestamp}.{format}"
        filepath = os.path.join(self.exports_dir, filename)
        
        if format == "pdf":
            self._export_to_pdf(filepath, session_name, messages, include_context)
        elif format == "docx":
            self._export_to_docx(filepath, session_name, messages, include_context)
        else:
            raise ValueError(f"Неподдерживаемый формат: {format}")
        
        return filepath
    
    def _export_to_pdf(
        self,
        filepath: str,
        session_name: str,
        messages: List[Dict],
        include_context: bool
    ):
        """
        Экспорт в PDF
        """
        doc = SimpleDocTemplate(filepath, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Заголовок
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor='#1a1a1a',
            spaceAfter=30,
            alignment=TA_LEFT
        )
        story.append(Paragraph(f"Сессия: {session_name}", title_style))
        story.append(Paragraph(f"Дата экспорта: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles['Normal']))
        story.append(Spacer(1, 0.3*inch))
        
        # Сообщения
        for i, msg in enumerate(messages, 1):
            # Промт
            if msg.get("prompt"):
                story.append(Paragraph(f"<b>Вопрос {i}:</b>", styles['Heading2']))
                story.append(Paragraph(msg["prompt"], styles['Normal']))
                story.append(Spacer(1, 0.1*inch))
            
            # Ответ
            if msg.get("response"):
                story.append(Paragraph(f"<b>Ответ YandexGPT:</b>", styles['Heading3']))
                story.append(Paragraph(msg["response"], styles['Normal']))
                story.append(Spacer(1, 0.2*inch))
            
            if include_context and msg.get("timestamp"):
                story.append(Paragraph(
                    f"<i>Время: {datetime.fromisoformat(msg['timestamp']).strftime('%d.%m.%Y %H:%M:%S')}</i>",
                    styles['Italic']
                ))
                story.append(Spacer(1, 0.2*inch))
            
            if i < len(messages):
                story.append(PageBreak())
        
        doc.build(story)
    
    def _export_to_docx(
        self,
        filepath: str,
        session_name: str,
        messages: List[Dict],
        include_context: bool
    ):
        """
        Экспорт в DOCX
        """
        doc = Document()
        
        # Заголовок
        title = doc.add_heading(f"Сессия: {session_name}", 0)
        doc.add_paragraph(f"Дата экспорта: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        doc.add_paragraph()
        
        # Сообщения
        for i, msg in enumerate(messages, 1):
            # Промт
            if msg.get("prompt"):
                doc.add_heading(f"Вопрос {i}", level=2)
                doc.add_paragraph(msg["prompt"])
            
            # Ответ
            if msg.get("response"):
                doc.add_heading("Ответ YandexGPT", level=3)
                doc.add_paragraph(msg["response"])
            
            if include_context and msg.get("timestamp"):
                p = doc.add_paragraph()
                p.add_run(f"Время: {datetime.fromisoformat(msg['timestamp']).strftime('%d.%m.%Y %H:%M:%S')}").italic = True
            
            doc.add_paragraph()
        
        doc.save(filepath)
